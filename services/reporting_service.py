"""
Reporting Service
Builds CSV, XLSX, and PDF exports from any module query result.
Depends on: openpyxl (XLSX), reportlab (PDF payslips).
"""
import csv
import io
from datetime import date, datetime

from flask import current_app
from services.db import get_cursor

# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------

def to_csv(rows, columns=None):
    """
    Convert a list of dicts (psycopg2 RealDictRow) to CSV bytes.
    `columns` controls the column order and can rename headers:
        [('employee_no', 'Employee No'), ('full_name', 'Full Name'), ...]
    If omitted, uses the dict keys from the first row.
    """
    if not rows:
        return b''

    buf = io.StringIO()
    writer = csv.writer(buf)

    if columns:
        keys = [c[0] for c in columns]
        headers = [c[1] for c in columns]
    else:
        keys = list(rows[0].keys())
        headers = keys

    writer.writerow(headers)
    for row in rows:
        writer.writerow([row.get(k, '') for k in keys])

    return buf.getvalue().encode('utf-8-sig')  # BOM for Excel compatibility


# ---------------------------------------------------------------------------
# XLSX export
# ---------------------------------------------------------------------------

def to_xlsx(rows, sheet_name='Report', columns=None, title=None):
    """
    Convert query results to an XLSX workbook (bytes).
    Requires: pip install openpyxl
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise RuntimeError("openpyxl is required for XLSX export: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # sheet name max 31 chars

    header_fill  = PatternFill('solid', fgColor='1F4E79')
    header_font  = Font(bold=True, color='FFFFFF')
    header_align = Alignment(horizontal='center')

    start_row = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
        start_row = 3

    if not rows:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    if columns:
        keys    = [c[0] for c in columns]
        headers = [c[1] for c in columns]
    else:
        keys    = list(rows[0].keys())
        headers = keys

    # Write header row
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill  = header_fill
        cell.font  = header_font
        cell.alignment = header_align

    # Write data rows
    for row_idx, row in enumerate(rows, start_row + 1):
        for col_idx, key in enumerate(keys, 1):
            val = row.get(key, '')
            # Convert date objects to string so Excel handles them consistently
            if isinstance(val, date):
                val = val.isoformat()
            ws.cell(row=row_idx, column=col_idx, value=val)

    # Auto-size columns (approximate)
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF payslip
# ---------------------------------------------------------------------------

def render_payslip_pdf(employee_id, pay_run_id):
    """
    Build a single payslip PDF using reportlab.
    Requires: pip install reportlab
    Returns bytes.
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                        Spacer, Paragraph)
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
    except ImportError:
        raise RuntimeError("reportlab is required for PDF payslips: pip install reportlab")

    # Fetch payslip data
    with get_cursor() as cur:
        cur.execute("""
            SELECT * FROM payroll.v_payslip
            WHERE employee_id = %s AND run_id = %s
        """, (employee_id, pay_run_id))
        payslip = cur.fetchone()

    if not payslip:
        raise ValueError(f"No payslip found for employee {employee_id} in run {pay_run_id}")

    # Fetch company branding
    with get_cursor() as cur:
        cur.execute("SELECT key, value FROM core.company_branding ORDER BY key")
        branding = {r['key']: r['value'] for r in cur.fetchall()}

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=15*mm, bottomMargin=15*mm,
                            leftMargin=20*mm, rightMargin=20*mm)
    styles = getSampleStyleSheet()
    bold   = ParagraphStyle('bold', parent=styles['Normal'], fontName='Helvetica-Bold')
    story  = []

    company_name = branding.get('company_name', 'Company')
    story.append(Paragraph(company_name, styles['Title']))
    story.append(Paragraph('PAYSLIP', bold))
    story.append(Spacer(1, 5*mm))

    # Employee info
    info = [
        ['Employee No:', payslip.get('employee_no', ''),
         'Period:', f"{payslip.get('period_start','')} – {payslip.get('period_end','')}"],
        ['Name:',        payslip.get('full_name', ''),
         'Pay Date:',    str(payslip.get('pay_date', ''))],
        ['Department:',  payslip.get('dept_name', ''),
         'Position:',    payslip.get('position_title', '')],
    ]
    info_table = Table(info, colWidths=[35*mm, 55*mm, 30*mm, 55*mm])
    info_table.setStyle(TableStyle([
        ('FONTNAME',  (0,0),(-1,-1), 'Helvetica'),
        ('FONTSIZE',  (0,0),(-1,-1), 9),
        ('FONTNAME',  (0,0),(0,-1), 'Helvetica-Bold'),
        ('FONTNAME',  (2,0),(2,-1), 'Helvetica-Bold'),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 5*mm))

    # Earnings / Deductions table
    def _fmt(val):
        if val is None:
            return '—'
        try:
            return f'{float(val):,.2f}'
        except (ValueError, TypeError):
            return str(val)

    earnings = [
        ['EARNINGS', 'AMOUNT', 'DEDUCTIONS', 'AMOUNT'],
        ['Basic Pay',         _fmt(payslip.get('basic_pay')),
         'SSS',               _fmt(payslip.get('sss_ee'))],
        ['OT Pay',            _fmt(payslip.get('ot_pay')),
         'PhilHealth',        _fmt(payslip.get('philhealth_ee'))],
        ['Holiday Pay',       _fmt(payslip.get('holiday_pay')),
         'HDMF/Pag-IBIG',     _fmt(payslip.get('hdmf_ee'))],
        ['Night Diff',        _fmt(payslip.get('night_diff_pay')),
         'Withholding Tax',   _fmt(payslip.get('withholding_tax'))],
        ['Allowances',        _fmt(payslip.get('total_allowances')),
         'Loan Deductions',   _fmt(payslip.get('total_loan_deductions'))],
        ['Other Earnings',    _fmt(payslip.get('other_earnings')),
         'Other Deductions',  _fmt(payslip.get('other_deductions'))],
        ['GROSS PAY',         _fmt(payslip.get('gross_pay')),
         'TOTAL DEDUCTIONS',  _fmt(payslip.get('total_deductions'))],
        ['', '', 'NET PAY',   _fmt(payslip.get('net_pay'))],
    ]
    pay_table = Table(earnings, colWidths=[55*mm, 30*mm, 55*mm, 35*mm])
    pay_table.setStyle(TableStyle([
        ('BACKGROUND',  (0,0), (-1,0), colors.HexColor('#1F4E79')),
        ('TEXTCOLOR',   (0,0), (-1,0), colors.white),
        ('FONTNAME',    (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTNAME',    (0,-2),(-1,-1),'Helvetica-Bold'),
        ('FONTSIZE',    (0,0), (-1,-1), 9),
        ('ALIGN',       (1,0), (1,-1), 'RIGHT'),
        ('ALIGN',       (3,0), (3,-1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#EBF0F8')]),
        ('GRID',        (0,0), (-1,-1), 0.5, colors.grey),
        ('TOPPADDING',  (0,0), (-1,-1), 3),
        ('BOTTOMPADDING',(0,0),(-1,-1), 3),
    ]))
    story.append(pay_table)
    story.append(Spacer(1, 8*mm))
    story.append(Paragraph(
        'This is a system-generated payslip. No signature required.',
        ParagraphStyle('footer', parent=styles['Normal'], fontSize=8,
                       textColor=colors.grey)))

    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Generic query-based report
# ---------------------------------------------------------------------------

def to_pdf(rows, columns=None, title=None, subtitle=None, landscape=False):
    """
    Convert query results to a PDF table report.
    Requires: pip install reportlab
    Returns bytes.
    """
    try:
        from reportlab.lib.pagesizes import A4, landscape as ls_mode
        from reportlab.lib import colors
        from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                        Spacer, Paragraph, PageBreak)
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
    except ImportError:
        raise RuntimeError("reportlab is required for PDF export: pip install reportlab")

    if not rows:
        rows = []

    if columns:
        keys = [c[0] for c in columns]
        headers = [c[1] for c in columns]
    elif rows:
        keys = list(rows[0].keys())
        headers = [k.replace('_', ' ').title() for k in keys]
    else:
        keys, headers = [], []

    buf = io.BytesIO()
    pagesize = ls_mode(A4) if landscape else A4
    doc = SimpleDocTemplate(buf, pagesize=pagesize,
                            topMargin=12*mm, bottomMargin=12*mm,
                            leftMargin=12*mm, rightMargin=12*mm)
    styles = getSampleStyleSheet()
    story = []

    # Header
    if title:
        story.append(Paragraph(title, styles['Title']))
    if subtitle:
        story.append(Paragraph(subtitle,
                     ParagraphStyle('sub', parent=styles['Normal'],
                                    fontSize=9, textColor=colors.grey)))
    story.append(Paragraph(
        f'Generated: {date.today().strftime("%B %d, %Y")} &bull; {len(rows)} records',
        ParagraphStyle('meta', parent=styles['Normal'],
                       fontSize=8, textColor=colors.grey)))
    story.append(Spacer(1, 4*mm))

    if not rows:
        story.append(Paragraph('No data found.', styles['Normal']))
        doc.build(story)
        return buf.getvalue()

    # Determine column widths
    page_w = pagesize[0] - 24*mm
    n_cols = len(headers)
    col_w = page_w / max(n_cols, 1)

    # Table data
    table_data = [headers]
    for row in rows[:2000]:  # cap at 2000 rows for PDF
        table_data.append([_format_cell(row.get(k, '')) for k in keys])

    # Build table with alternating rows
    t = Table(table_data, colWidths=[col_w] * n_cols, repeatRows=1)
    style_cmds = [
        ('BACKGROUND',    (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
        ('TEXTCOLOR',     (0, 0), (-1, 0), colors.white),
        ('FONTNAME',      (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE',      (0, 0), (-1, 0), 8),
        ('FONTSIZE',      (0, 1), (-1, -1), 7),
        ('FONTNAME',      (0, 1), (-1, -1), 'Helvetica'),
        ('ALIGN',         (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING',    (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('GRID',          (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc')),
        ('ROWBACKGROUNDS',(0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F4F8')]),
    ]
    # Right-align numeric columns
    for ci, k in enumerate(keys):
        if rows and _is_numeric(rows[0].get(k)):
            style_cmds.append(('ALIGN', (ci, 0), (ci, -1), 'RIGHT'))
    t.setStyle(TableStyle(style_cmds))
    story.append(t)

    # Footer
    story.append(Spacer(1, 6*mm))
    story.append(Paragraph(
        'HCM360 HRIS — System-generated report. Confidential.',
        ParagraphStyle('footer', parent=styles['Normal'],
                       fontSize=7, textColor=colors.grey)))

    doc.build(story)
    return buf.getvalue()


def _format_cell(val):
    """Format cell values for PDF display."""
    if val is None:
        return ''
    if isinstance(val, date):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d %H:%M')
    try:
        fv = float(val)
        if fv == int(fv) and abs(fv) < 1e10:
            return str(int(fv))
        return f'{fv:,.2f}'
    except (ValueError, TypeError):
        return str(val)[:80]


def _is_numeric(val):
    if val is None:
        return False
    try:
        float(val)
        return True
    except (ValueError, TypeError):
        return False


def build_report(sql, params=None, fmt='csv', sheet_name='Report',
                 columns=None, title=None):
    """
    Execute `sql` with `params` and export results.
    fmt: 'csv' | 'xlsx' | 'pdf'
    Returns (bytes, mimetype, filename_extension).
    """
    with get_cursor() as cur:
        cur.execute(sql, params or ())
        rows = cur.fetchall()

    if fmt == 'xlsx':
        data = to_xlsx(rows, sheet_name=sheet_name, columns=columns, title=title)
        return data, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'
    elif fmt == 'pdf':
        data = to_pdf(rows, columns=columns, title=title,
                      landscape=len(rows[0].keys() if rows else []) > 6)
        return data, 'application/pdf', 'pdf'
    else:
        data = to_csv(rows, columns=columns)
        return data, 'text/csv', 'csv'


# ---------------------------------------------------------------------------
# Pre-built module report queries
# ---------------------------------------------------------------------------

REPORT_QUERIES = {
    'headcount': """
        SELECT department, position_title, employment_type,
               COUNT(*) AS headcount
        FROM core.v_employees_full
        WHERE employment_status = 'ACTIVE'
        GROUP BY department, position_title, employment_type
        ORDER BY department, headcount DESC
    """,
    'attendance_summary': """
        SELECT
            e.employee_no, e.full_name, e.department,
            COUNT(ad.work_date)                       AS days_worked,
            SUM(CASE WHEN ad.is_late THEN 1 ELSE 0 END)     AS late_count,
            SUM(CASE WHEN ad.is_absent THEN 1 ELSE 0 END)   AS absent_count,
            ROUND(AVG(ad.late_minutes)::NUMERIC, 1)          AS avg_late_min,
            SUM(ad.ot_hours)                                  AS total_ot_hours
        FROM core.v_employees_full e
        LEFT JOIN attendance.att_daily ad ON ad.employee_id = e.id
          AND ad.work_date BETWEEN %(date_from)s AND %(date_to)s
        GROUP BY e.employee_no, e.full_name, e.department
        ORDER BY e.department, e.full_name
    """,
    'leave_utilization': """
        SELECT
            e.employee_no, e.full_name, e.department,
            lt.code AS leave_type, lt.name AS leave_type_name,
            COUNT(lr.id)      AS requests,
            SUM(lr.total_days) AS total_days,
            lb.entitled_days, lb.used_days, lb.balance
        FROM core.v_employees_full e
        JOIN leave_mgmt.lv_balances lb ON lb.employee_id = e.id
        JOIN leave_mgmt.lv_types lt    ON lt.id = lb.leave_type_id
        LEFT JOIN leave_mgmt.lv_requests lr ON lr.employee_id = e.id
          AND lr.leave_type_id = lt.id AND lr.status = 'APPROVED'
        WHERE e.is_active = TRUE
        GROUP BY e.employee_no, e.full_name, e.department,
                 lt.code, lt.name, lb.entitled_days, lb.used_days, lb.balance
        ORDER BY e.department, e.full_name, lt.code
    """,
    'payroll_register': """
        SELECT
            e.employee_no, e.full_name, d.name AS department,
            pp.basic_pay, pp.ot_pay, pp.gross_pay,
            pp.total_deductions, pp.net_pay,
            pr.period_start, pr.period_end, pr.pay_date
        FROM payroll.pay_employee_payroll pp
        JOIN core.v_employees_full e ON e.id = pp.employee_id
        JOIN core.departments d ON d.id = e.department_id
        JOIN payroll.pay_runs pr ON pr.id = pp.pay_run_id
        WHERE pp.pay_run_id = %(pay_run_id)s
        ORDER BY d.name, e.full_name
    """,
}
