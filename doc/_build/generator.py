#!/usr/bin/env python3
"""
HCM360 Documentation Generator

Produces .docx + .pdf from a shared content model.
Content is expressed as a list of blocks: (kind, *args).
"""
import os
import sys
from pathlib import Path

# ── DOCX imports ────────────────────────────────────────────────────
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── PDF imports ─────────────────────────────────────────────────────
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, PageBreak,
                                Table, TableStyle, ListFlowable, ListItem,
                                KeepTogether)

BRAND_BLUE = RGBColor(0x1E, 0x40, 0xAF)
BRAND_BLUE_HEX = colors.HexColor('#1E40AF')
INDIGO = colors.HexColor('#4F46E5')
GRAY = colors.HexColor('#6B7280')
LIGHT_GRAY = colors.HexColor('#F3F4F6')
DARK = colors.HexColor('#111827')

OUT_DIR = Path(__file__).parent.parent
OUT_DOCX = OUT_DIR
OUT_PDF = OUT_DIR


# ════════════════════════════════════════════════════════════════════
# DOCX RENDERER
# ════════════════════════════════════════════════════════════════════
def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _set_cell_margin(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def render_docx(title, subtitle, blocks, out_path):
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('HCM360 HRIS')
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = BRAND_BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        run.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Version 1.0.0  ·  April 2026')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    doc.add_paragraph()
    doc.add_paragraph()

    for block in blocks:
        kind = block[0]

        if kind == 'H1':
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = BRAND_BLUE
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)

        elif kind == 'H2':
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            run.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)

        elif kind == 'H3':
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)

        elif kind == 'P':
            p = doc.add_paragraph(block[1])
            p.paragraph_format.space_after = Pt(6)

        elif kind == 'UL':
            for item in block[1]:
                p = doc.add_paragraph(item, style='List Bullet')
                p.paragraph_format.space_after = Pt(2)

        elif kind == 'OL':
            for item in block[1]:
                p = doc.add_paragraph(item, style='List Number')
                p.paragraph_format.space_after = Pt(2)

        elif kind == 'TABLE':
            headers = block[1]
            rows = block[2]
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = 'Light Grid Accent 1'
            # Header row
            for i, h in enumerate(headers):
                cell = t.rows[0].cells[i]
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(h)
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(cell, '1E40AF')
                _set_cell_margin(cell)
            # Data rows
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    cell = t.rows[r_idx + 1].cells[c_idx]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    run = p.add_run(str(val))
                    run.font.size = Pt(10)
                    _set_cell_margin(cell)
                    if r_idx % 2 == 1:
                        _shade(cell, 'F8FAFC')
            # Spacing
            doc.add_paragraph()

        elif kind == 'CODE':
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(6)

        elif kind == 'NOTE':
            p = doc.add_paragraph()
            run = p.add_run('Note: ')
            run.bold = True
            run.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)
            run2 = p.add_run(block[1])
            run2.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)

        elif kind == 'PB':
            doc.add_page_break()

    doc.save(str(out_path))
    return out_path


# ════════════════════════════════════════════════════════════════════
# PDF RENDERER
# ════════════════════════════════════════════════════════════════════
def _make_pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle('CoverBrand', fontName='Helvetica-Bold', fontSize=14,
                              alignment=TA_CENTER, textColor=BRAND_BLUE_HEX, spaceAfter=6))
    styles.add(ParagraphStyle('CoverTitle', fontName='Helvetica-Bold', fontSize=26,
                              alignment=TA_CENTER, textColor=DARK, spaceAfter=12))
    styles.add(ParagraphStyle('CoverSubtitle', fontName='Helvetica-Oblique', fontSize=13,
                              alignment=TA_CENTER, textColor=GRAY, spaceAfter=24))
    styles.add(ParagraphStyle('CoverMeta', fontName='Helvetica', fontSize=10,
                              alignment=TA_CENTER, textColor=colors.HexColor('#9CA3AF')))
    styles.add(ParagraphStyle('MyH1', fontName='Helvetica-Bold', fontSize=18,
                              textColor=BRAND_BLUE_HEX, spaceBefore=18, spaceAfter=8,
                              keepWithNext=1))
    styles.add(ParagraphStyle('MyH2', fontName='Helvetica-Bold', fontSize=14,
                              textColor=DARK, spaceBefore=12, spaceAfter=6,
                              keepWithNext=1))
    styles.add(ParagraphStyle('MyH3', fontName='Helvetica-Bold', fontSize=11,
                              textColor=colors.HexColor('#374151'),
                              spaceBefore=8, spaceAfter=4, keepWithNext=1))
    styles.add(ParagraphStyle('MyBody', fontName='Helvetica', fontSize=10.5, leading=15,
                              textColor=DARK, spaceAfter=5, alignment=TA_LEFT))
    styles.add(ParagraphStyle('MyCode', fontName='Courier', fontSize=8.5, leading=11,
                              textColor=colors.HexColor('#1F2937'),
                              backColor=colors.HexColor('#F3F4F6'),
                              borderPadding=6, leftIndent=12, rightIndent=12,
                              spaceBefore=4, spaceAfter=8))
    styles.add(ParagraphStyle('MyNote', fontName='Helvetica', fontSize=10, leading=14,
                              textColor=colors.HexColor('#065F46'),
                              backColor=colors.HexColor('#F0FDF4'),
                              borderColor=colors.HexColor('#86EFAC'), borderWidth=1,
                              borderPadding=8, leftIndent=6, rightIndent=6,
                              spaceBefore=6, spaceAfter=8))
    return styles


def _escape(text):
    if not isinstance(text, str):
        text = str(text)
    return (text.replace('&', '&amp;')
                .replace('<', '&lt;')
                .replace('>', '&gt;'))


def render_pdf(title, subtitle, blocks, out_path):
    styles = _make_pdf_styles()
    doc = SimpleDocTemplate(str(out_path), pagesize=LETTER,
                            leftMargin=0.85 * inch, rightMargin=0.85 * inch,
                            topMargin=0.85 * inch, bottomMargin=0.85 * inch,
                            title=title, author='HCM360')
    story = []

    # Cover
    story.append(Spacer(1, 1.2 * inch))
    story.append(Paragraph('HCM360 HRIS', styles['CoverBrand']))
    story.append(Paragraph(_escape(title), styles['CoverTitle']))
    if subtitle:
        story.append(Paragraph(_escape(subtitle), styles['CoverSubtitle']))
    story.append(Spacer(1, 0.5 * inch))
    story.append(Paragraph('Version 1.0.0 &nbsp;·&nbsp; April 2026', styles['CoverMeta']))
    story.append(PageBreak())

    for block in blocks:
        kind = block[0]

        if kind == 'H1':
            story.append(Paragraph(_escape(block[1]), styles['MyH1']))

        elif kind == 'H2':
            story.append(Paragraph(_escape(block[1]), styles['MyH2']))

        elif kind == 'H3':
            story.append(Paragraph(_escape(block[1]), styles['MyH3']))

        elif kind == 'P':
            story.append(Paragraph(_escape(block[1]), styles['MyBody']))

        elif kind == 'UL':
            items = [ListItem(Paragraph(_escape(i), styles['MyBody']),
                              leftIndent=18, bulletColor=INDIGO)
                     for i in block[1]]
            story.append(ListFlowable(items, bulletType='bullet', start='•',
                                      leftIndent=20, spaceBefore=2, spaceAfter=8))

        elif kind == 'OL':
            items = [ListItem(Paragraph(_escape(i), styles['MyBody']),
                              leftIndent=20)
                     for i in block[1]]
            story.append(ListFlowable(items, bulletType='1', leftIndent=24,
                                      spaceBefore=2, spaceAfter=8))

        elif kind == 'TABLE':
            headers = block[1]
            rows = block[2]
            all_rows = [headers] + [[str(c) for c in r] for r in rows]
            # Wrap each cell in Paragraph for wrapping
            pstyle = ParagraphStyle('tbl', fontName='Helvetica', fontSize=9,
                                    leading=12, textColor=DARK)
            hstyle = ParagraphStyle('tblh', fontName='Helvetica-Bold', fontSize=9,
                                    leading=12, textColor=colors.white)
            wrapped = [[Paragraph(_escape(c), hstyle) for c in all_rows[0]]]
            for r in all_rows[1:]:
                wrapped.append([Paragraph(_escape(c), pstyle) for c in r])

            # Width: split evenly across available space (~7 inches)
            n = len(headers)
            avail = 7.0 * inch
            col_w = [avail / n] * n
            t = Table(wrapped, colWidths=col_w, repeatRows=1)
            ts = [
                ('BACKGROUND', (0, 0), (-1, 0), BRAND_BLUE_HEX),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]
            # Zebra stripes
            for r_idx in range(1, len(wrapped)):
                if r_idx % 2 == 0:
                    ts.append(('BACKGROUND', (0, r_idx), (-1, r_idx),
                               colors.HexColor('#F9FAFB')))
            t.setStyle(TableStyle(ts))
            story.append(t)
            story.append(Spacer(1, 0.15 * inch))

        elif kind == 'CODE':
            text = _escape(block[1]).replace('\n', '<br/>').replace(' ', '&nbsp;')
            story.append(Paragraph(f'<font face="Courier">{text}</font>',
                                   styles['MyCode']))

        elif kind == 'NOTE':
            story.append(Paragraph('<b>Note:</b> ' + _escape(block[1]),
                                   styles['MyNote']))

        elif kind == 'PB':
            story.append(PageBreak())

    def _footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(GRAY)
        canvas.drawString(0.85 * inch, 0.5 * inch,
                          f'HCM360 HRIS · {title}')
        canvas.drawRightString(LETTER[0] - 0.85 * inch, 0.5 * inch,
                               f'Page {doc_.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return out_path


# ════════════════════════════════════════════════════════════════════
# DOC INDEX
# ════════════════════════════════════════════════════════════════════
# Imports each doc module — they each expose TITLE, SUBTITLE, FOLDER, BLOCKS
DOCS = [
    '01_architecture_overview',
    '02_data_structure',
    '03_apis_and_logic',
    '04_deployment_setup',
    '05_access_and_security',
    '06_user_guide',
    '07_process_flows',
    '08_feature_overview',
    '09_admin_guide',
    '10_kpi_explanation',
    '11_competitive_advantage',
]


def main():
    import runpy
    content_dir = Path(__file__).parent / 'content'
    results = []
    for mod_name in DOCS:
        path = content_dir / f'{mod_name}.py'
        ns = runpy.run_path(str(path))
        title = ns['TITLE']
        subtitle = ns['SUBTITLE']
        folder = ns['FOLDER']
        blocks = ns['BLOCKS']
        out_dir = OUT_DIR / folder
        out_dir.mkdir(parents=True, exist_ok=True)
        stem = mod_name
        docx_path = out_dir / f'{stem}.docx'
        pdf_path = out_dir / f'{stem}.pdf'
        render_docx(title, subtitle, blocks, docx_path)
        render_pdf(title, subtitle, blocks, pdf_path)
        size_docx = docx_path.stat().st_size
        size_pdf = pdf_path.stat().st_size
        results.append((mod_name, title, size_docx, size_pdf))
        print(f'✓ {mod_name}: docx={size_docx // 1024}KB, pdf={size_pdf // 1024}KB')
    return results


if __name__ == '__main__':
    main()
